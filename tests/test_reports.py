import io

from openpyxl import load_workbook

from app.reports.builder import build_summary_report_xlsx, build_test_report_csv, build_test_report_xlsx
from app.services import attempts as attempt_service
from app.services import test_builder as tb
from tests.factories import make_employee, make_material_with_text, make_question, make_test


async def _setup(session_maker, session):
    mid = await make_material_with_text(session_maker)
    for i in range(3):
        await make_question(session, text_=f"Hisobot savoli {i} nima haqida?", material_id=mid, correct="A")
    test = await make_test(session, question_count=3, randomize_options=False)
    await tb.assemble_questions(session_maker, test.id, None)
    await tb.mark_ready(session, test.id)
    await tb.publish(session, test.id)
    done = await make_employee(session, 300, "Hisobot Bajaruvchi")
    await make_employee(session, 301, "Hisobot Qatnashmagan")
    attempt = (await attempt_service.start_attempt(session, done, test.id)).attempt
    for pos, letter in enumerate("AAB"):
        await session.refresh(attempt)
        await attempt_service.submit_answer(session, done, attempt.id, pos, letter)
    return test


async def test_test_report_xlsx_contents(session_maker, session):
    test = await _setup(session_maker, session)
    data, filename = await build_test_report_xlsx(session, test.id)
    assert filename.endswith(".xlsx")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == [
        "Umumiy",
        "Ishtirokchilar",
        "Qatnashmaganlar",
        "Barcha urinishlar",
        "Xato javoblar",
        "Barcha javoblar",
        "Savollar tahlili",
        "Manba bo'yicha",
    ]
    attempts = list(wb["Barcha urinishlar"].iter_rows(values_only=True))
    header, row = attempts[0], attempts[1]
    assert header[0] == "Xodim" and "Foiz" in header
    record = dict(zip(header, row, strict=True))
    assert record["Xodim"] == "Hisobot Bajaruvchi"
    assert record["Telegram ID"] == 300
    assert record["Jami savollar"] == 3 and record["To'g'ri"] == 2 and record["Noto'g'ri"] == 1
    assert abs(record["Foiz"] - 66.67) < 0.01 and record["Holat"] == "Yakunlagan"
    not_participated = [r[0] for r in wb["Qatnashmaganlar"].iter_rows(min_row=2, values_only=True)]
    assert not_participated == ["Hisobot Qatnashmagan"]
    every = list(wb["Barcha javoblar"].iter_rows(min_row=2, values_only=True))
    assert len(every) == 3 and [r[8] for r in every].count("Xato") == 1
    wrong = list(wb["Xato javoblar"].iter_rows(min_row=2, values_only=True))
    assert len(wrong) == 1 and wrong[0][4].startswith("B)") and wrong[0][5].startswith("A)")


async def test_summary_report_and_csv(session_maker, session):
    test = await _setup(session_maker, session)
    data, _ = await build_summary_report_xlsx(session)
    wb = load_workbook(io.BytesIO(data))
    assert "Xodimlar bo'yicha" in wb.sheetnames
    summary = list(wb["Xodimlar bo'yicha"].iter_rows(min_row=2, values_only=True))
    assert summary[0][0] == "Hisobot Bajaruvchi" and summary[0][7] == 2
    csv_bytes, name = await build_test_report_csv(session, test.id)
    text = csv_bytes.decode("utf-8-sig")
    assert name.endswith(".csv") and "Hisobot Bajaruvchi" in text and text.count("\n") == 2
