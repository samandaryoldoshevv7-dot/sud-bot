import itertools
from pathlib import Path

import pytest

from app.documents import DocumentError, chunk_pages, clean_pages, detect_file_type, extract_document
from app.documents.extractors import ExtractedPage
from app.models import FileType
from tests.factories import LAW_TEXT


def test_detect_file_type():
    assert detect_file_type("a.PDF", None) == FileType.PDF
    assert detect_file_type("a.docx", None) == FileType.DOCX
    assert detect_file_type("notes.md", None) == FileType.MD
    assert detect_file_type("x.txt", "text/plain") == FileType.TXT
    with pytest.raises(DocumentError) as exc:
        detect_file_type("virus.exe", "application/octet-stream")
    assert exc.value.code == "unsupported_type"
    with pytest.raises(DocumentError) as exc:
        detect_file_type("old.doc", None)
    assert exc.value.code == "unsupported_doc"


def test_txt_and_markdown_extraction(tmp_path: Path):
    p = tmp_path / "a.md"
    p.write_text("# Sarlavha\n\n**Muhim** qoida: [havola](http://x) matni.\n\n- band bir\n", encoding="utf-8")
    doc = extract_document(p, FileType.MD, 10_000)
    assert "Muhim qoida: havola matni." in doc.text
    assert "**" not in doc.text and "http" not in doc.text

    p2 = tmp_path / "b.txt"
    p2.write_bytes("Кирилл матн".encode("cp1251"))
    assert "Кирилл" in extract_document(p2, FileType.TXT, 10_000).text


def test_docx_extraction_with_table(tmp_path: Path):
    from docx import Document

    d = Document()
    d.add_heading("1-modda. Umumiy qoidalar", level=1)
    d.add_paragraph("Ariza kelib tushgan kunning o'zida ro'yxatga olinadi.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Muddat"
    table.rows[0].cells[1].text = "o'n kun"
    path = tmp_path / "x.docx"
    d.save(path)
    doc = extract_document(path, FileType.DOCX, 10_000)
    assert "1-modda. Umumiy qoidalar" in doc.text
    assert "ro'yxatga olinadi" in doc.text
    assert "Muddat | o'n kun" in doc.text


def test_pdf_extraction_keeps_page_numbers(tmp_path: Path):
    from fpdf import FPDF

    pdf = FPDF()
    for i in range(1, 4):
        pdf.add_page()
        pdf.set_font("helvetica", size=12)
        pdf.multi_cell(0, 10, f"Page {i} content. The appeal must be filed within ten days on page {i}.")
    path = tmp_path / "x.pdf"
    pdf.output(str(path))
    doc = extract_document(path, FileType.PDF, 100_000)
    assert doc.page_count == 3
    assert [p.page_no for p in doc.pages] == [1, 2, 3]
    assert "within ten days on page 2" in doc.pages[1].text


def test_pdf_without_text_is_rejected(tmp_path: Path):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    path = tmp_path / "empty.pdf"
    pdf.output(str(path))
    with pytest.raises(DocumentError) as exc:
        extract_document(path, FileType.PDF, 100_000)
    assert exc.value.code == "pdf_no_text"


def test_corrupted_pdf(tmp_path: Path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"%PDF-1.4 garbage")
    with pytest.raises(DocumentError) as exc:
        extract_document(path, FileType.PDF, 100_000)
    assert exc.value.code in ("pdf_corrupted", "pdf_no_text")


def test_text_size_limit(tmp_path: Path):
    path = tmp_path / "big.txt"
    path.write_text("a" * 5000)
    with pytest.raises(DocumentError) as exc:
        extract_document(path, FileType.TXT, 1000)
    assert exc.value.code == "too_large_text"


def test_cleaner_removes_headers_and_page_numbers():
    pages = [ExtractedPage(i, f"OLIY SUD AXBOROTNOMASI\nMatn {i} qismi davom-\netadi.\n{i}") for i in range(1, 6)]
    cleaned = clean_pages(pages)
    for page in cleaned:
        assert "AXBOROTNOMASI" not in page.text
        assert "davometadi" in page.text
        assert not page.text.strip().endswith(str(page.page_no))


def test_chunker_respects_size_overlap_and_pages():
    pages = [ExtractedPage(1, LAW_TEXT), ExtractedPage(2, LAW_TEXT)]
    chunks = chunk_pages(clean_pages(pages), chunk_size=500, overlap=120, min_chars=100)
    assert len(chunks) >= 4
    assert all(c.char_count <= 650 for c in chunks)
    assert chunks[0].page_start == 1 and chunks[-1].page_end == 2
    assert any(c.heading and "modda" in c.heading for c in chunks)
    # Overlap: consecutive chunks share some trailing text.
    shared = sum(1 for a, b in itertools.pairwise(chunks) if a.text.split("\n\n")[-1] in b.text)
    assert shared >= 1
