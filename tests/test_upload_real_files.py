"""Uploading real PDF and DOCX files through the bot (both upload flows), as an admin does it."""

import asyncio
import io

import docx
from fpdf import FPDF
from sqlalchemy import select

from app.keyboards.callbacks import AdminCB
from app.models import Material, MaterialStatus
from tests.factories import make_employee
from tests.telegram_mock import callback_update, document_update, message_update
from tests.test_group_flow import tg  # noqa: F401  (fixture)

ADMIN = 1000
PARAGRAPHS = [
    "1-modda. Sudga kelib tushgan har bir ariza kelib tushgan kunning o'zida ro'yxatga olinadi.",
    "2-modda. Apellyatsiya shikoyati hal qiluv qarori e'lon qilingan kundan e'tiboran o'n kun ichida beriladi.",
    "3-modda. Sud majlisining bayonnomasi majlis tugaganidan keyin uch kun ichida rasmiylashtiriladi.",
    "4-modda. Ko'rib chiqilgan ishlar sud qarori qonuniy kuchga kirganidan so'ng bir oy ichida arxivga topshiriladi.",
] * 3


def _pdf() -> bytes:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for _ in range(2):
        pdf.add_page()
        for p in PARAGRAPHS:
            pdf.multi_cell(w=180, h=6, text=p, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _docx() -> bytes:
    document = docx.Document()
    document.add_heading("Sud ishlarini yuritish", 1)
    for p in PARAGRAPHS:
        document.add_paragraph(p)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


async def _wait():
    from app.services import background

    for _ in range(200):
        if background.running_count() == 0:
            return
        await asyncio.sleep(0.05)


FILES = [
    ("pdf1", "Yo'riqnoma.pdf", "application/pdf", _pdf),
    ("docx1", "Mehnat kodeksi.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", _docx),
]


async def test_materials_section_accepts_pdf_and_docx(tg, session_maker):  # noqa: F811
    async with session_maker() as session:
        await make_employee(session, ADMIN, "Admin Bosh")
    for file_id, name, mime, build in FILES:
        content = build()
        tg.session.files[file_id] = content
        tg.session.clear()
        await tg(callback_update(ADMIN, AdminCB(s="mat_up").pack()))
        await tg(document_update(ADMIN, file_id, name, len(content), mime))
        await tg(callback_update(ADMIN, AdminCB(s="mat_t_def").pack()))
        await tg(callback_update(ADMIN, AdminCB(s="mat_cat", v="").pack()))
        await tg(callback_update(ADMIN, AdminCB(s="mat_desc_skip").pack()))
        await _wait()
        texts = "\n".join(tg.session.texts())
        assert "xatolik" not in texts.lower(), texts
    async with session_maker() as session:
        materials = (await session.execute(select(Material))).scalars().all()
        assert [m.status for m in materials] == [MaterialStatus.READY, MaterialStatus.READY], [
            (m.file_name, m.status, m.error_message) for m in materials
        ]


async def test_create_test_flow_accepts_pdf_and_docx(tg, session_maker):  # noqa: F811
    async with session_maker() as session:
        await make_employee(session, ADMIN, "Admin Bosh")
    await tg(callback_update(ADMIN, AdminCB(s="ct").pack()))
    await tg(message_update(ADMIN, "Haqiqiy fayllar testi"))
    for file_id, name, mime, build in FILES:
        content = build()
        tg.session.files[file_id] = content
        await tg(document_update(ADMIN, file_id, name, len(content), mime))
    await _wait()
    texts = "\n".join(tg.session.texts())
    assert "xatolik" not in texts.lower(), texts
    async with session_maker() as session:
        materials = (await session.execute(select(Material))).scalars().all()
        assert all(m.status == MaterialStatus.READY for m in materials) and len(materials) == 2, [
            (m.file_name, m.status, m.error_message) for m in materials
        ]
