"""Text extraction for PDF, DOCX, TXT and Markdown files.

Extraction works from a file on disk (a temporary download) and processes PDFs page by
page so large documents are not duplicated in memory unnecessarily.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from app.models.enums import FileType

logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {".pdf": FileType.PDF, ".docx": FileType.DOCX, ".txt": FileType.TXT, ".md": FileType.MD,
                      ".markdown": FileType.MD}  # fmt: skip
ALLOWED_MIME = {
    "application/pdf": FileType.PDF,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": FileType.DOCX,
    "text/plain": FileType.TXT,
    "text/markdown": FileType.MD,
    "text/x-markdown": FileType.MD,
}


class DocumentError(Exception):
    """A user-facing document problem. ``code`` maps to a localized message."""

    def __init__(self, code: str, detail: str = ""):
        super().__init__(f"{code}: {detail}" if detail else code)
        self.code = code
        self.detail = detail


@dataclass
class ExtractedPage:
    page_no: int | None
    text: str


@dataclass
class ExtractedDocument:
    pages: list[ExtractedPage] = field(default_factory=list)
    page_count: int | None = None

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)


def detect_file_type(file_name: str | None, mime_type: str | None) -> FileType:
    suffix = Path(file_name or "").suffix.lower()
    if suffix in ALLOWED_EXTENSIONS:
        return ALLOWED_EXTENSIONS[suffix]
    if suffix in (".doc", ".rtf", ".odt"):
        raise DocumentError("unsupported_doc")
    if mime_type in ALLOWED_MIME and not suffix:
        return ALLOWED_MIME[mime_type]
    raise DocumentError("unsupported_type", suffix or (mime_type or ""))


def extract_document(path: Path, file_type: FileType, max_chars: int) -> ExtractedDocument:
    """Synchronous extraction (run it in a worker thread)."""
    if file_type == FileType.PDF:
        doc = _extract_pdf(path, max_chars)
    elif file_type == FileType.DOCX:
        doc = _extract_docx(path, max_chars)
    elif file_type in (FileType.TXT, FileType.MD, FileType.TEXT):
        raw = path.read_bytes()
        doc = extract_plain_text(_decode(raw), markdown=file_type == FileType.MD)
    else:  # pragma: no cover
        raise DocumentError("unsupported_type", str(file_type))
    if doc.char_count > max_chars:
        raise DocumentError("too_large_text", str(doc.char_count))
    return doc


def extract_plain_text(text: str, markdown: bool = False) -> ExtractedDocument:
    if markdown:
        text = _strip_markdown(text)
    return ExtractedDocument(pages=[ExtractedPage(page_no=None, text=text)], page_count=None)


def _decode(raw: bytes) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        raw = raw[3:]
    for encoding in ("utf-8", "utf-16", "cp1251"):
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding == "utf-16" and "\x00" in text:
            continue
        return text
    return raw.decode("utf-8", errors="replace")


def _extract_pdf(path: Path, max_chars: int) -> ExtractedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise DocumentError("pdf_encrypted") from exc
        pages: list[ExtractedPage] = []
        total = 0
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:  # a single broken page should not fail the document
                logger.warning("PDF page extraction failed", extra={"page": index, "error": str(exc)[:200]})
                text = ""
            total += len(text)
            if total > max_chars:
                raise DocumentError("too_large_text", str(total))
            pages.append(ExtractedPage(page_no=index, text=text))
        doc = ExtractedDocument(pages=pages, page_count=len(reader.pages))
    except DocumentError:
        raise
    except PdfReadError as exc:
        raise DocumentError("pdf_corrupted", str(exc)[:200]) from exc
    except Exception as exc:
        raise DocumentError("pdf_corrupted", str(exc)[:200]) from exc
    if doc.char_count < 20:
        raise DocumentError("pdf_no_text")
    return doc


_HEADING_STYLE = re.compile(r"^(heading|title|заголовок|sarlavha)", re.IGNORECASE)


def _extract_docx(path: Path, max_chars: int) -> ExtractedDocument:
    import zipfile

    from docx import Document
    from docx.oxml.ns import qn

    try:
        document = Document(str(path))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise DocumentError("docx_corrupted", str(exc)[:200]) from exc
    except Exception as exc:
        raise DocumentError("docx_corrupted", str(exc)[:200]) from exc

    blocks: list[str] = []
    total = 0
    body = document.element.body
    # Walk the body in document order so tables stay where they appear.
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
            if not text:
                continue
            style_el = child.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
            style = style_el.get(qn("w:val")) if style_el is not None else ""
            if style and _HEADING_STYLE.match(style):
                text = f"\n{text}\n"
            blocks.append(text)
        elif child.tag == qn("w:tbl"):
            for row in child.iter(qn("w:tr")):
                cells = []
                for cell in row.iter(qn("w:tc")):
                    cell_text = " ".join(
                        "".join(n.text or "" for n in p.iter(qn("w:t"))).strip() for p in cell.iter(qn("w:p"))
                    ).strip()
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    blocks.append(" | ".join(cells))
        else:
            continue
        total += len(blocks[-1]) if blocks else 0
        if total > max_chars:
            raise DocumentError("too_large_text", str(total))
    text = "\n\n".join(blocks)
    if len(text.strip()) < 20:
        raise DocumentError("empty_document")
    return ExtractedDocument(pages=[ExtractedPage(page_no=None, text=text)], page_count=None)


_MD_PATTERNS = [
    (re.compile(r"```.*?```", re.DOTALL), lambda m: m.group(0).strip("`")),
    (re.compile(r"!\[([^\]]*)\]\([^)]*\)"), r"\1"),
    (re.compile(r"\[([^\]]+)\]\([^)]*\)"), r"\1"),
    (re.compile(r"^\s{0,3}#{1,6}\s*", re.MULTILINE), "\n"),
    (re.compile(r"^\s*>\s?", re.MULTILINE), ""),
    (re.compile(r"(\*\*|__)(.+?)\1"), r"\2"),
    (re.compile(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])"), r"\1"),
    (re.compile(r"`([^`]+)`"), r"\1"),
    (re.compile(r"^\s*([-*_]\s*){3,}$", re.MULTILINE), ""),
    (re.compile(r"^\s*[-*+]\s+", re.MULTILINE), "• "),
]


def _strip_markdown(text: str) -> str:
    for pattern, repl in _MD_PATTERNS:
        text = pattern.sub(repl, text)
    return text
