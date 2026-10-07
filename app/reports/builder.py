"""XLSX / CSV report generation (in memory, returned as bytes for Telegram upload)."""

from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    AttemptStatus,
    ParticipationStatus,
    Test,
    TestAttempt,
    TestQuestion,
    TestStatus,
    User,
    UserAnswer,
    UserRole,
)
from app.statistics.dashboard import test_question_stats
from app.statistics.participation import participation
from app.utils.time import fmt_dt, fmt_duration, to_local, utcnow

STATUS_LABELS = {
    ParticipationStatus.COMPLETED: "Yakunlagan",
    ParticipationStatus.IN_PROGRESS: "Jarayonda",
    ParticipationStatus.EXPIRED: "Muddati o'tgan (tugatmagan)",
    ParticipationStatus.CANCELLED: "Bekor qilingan",
    ParticipationStatus.NOT_STARTED: "Qatnashmagan",
}
ATTEMPT_LABELS = {
    AttemptStatus.COMPLETED: "Yakunlagan",
    AttemptStatus.IN_PROGRESS: "Jarayonda",
    AttemptStatus.EXPIRED: "Muddati o'tgan",
    AttemptStatus.CANCELLED: "Bekor qilingan",
}

ATTEMPT_HEADERS = [
    "Xodim",
    "Telegram ID",
    "Username",
    "Test",
    "Urinish №",
    "Boshlangan",
    "Yakunlangan",
    "Davomiylik",
    "Davomiylik (sek)",
    "Jami savollar",
    "Javob berilgan",
    "To'g'ri",
    "Noto'g'ri",
    "Foiz",
    "O'tdi",
    "Holat",
]

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_sheet(ws: Worksheet, widths: dict[int, int] | None = None) -> None:
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column:
        ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        if widths and col in widths:
            ws.column_dimensions[letter].width = widths[col]
            continue
        longest = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(10, longest + 2), 60)


def _naive_local(dt: datetime | None) -> datetime | None:
    local = to_local(dt)
    return local.replace(tzinfo=None) if local else None


def _attempt_row(user: User, test_title: str, a: TestAttempt) -> list:
    return [
        user.display_name,
        user.telegram_id,
        f"@{user.username}" if user.username else "",
        test_title,
        a.attempt_no,
        _naive_local(a.started_at),
        _naive_local(a.completed_at),
        fmt_duration(a.duration_seconds) if a.duration_seconds is not None else "",
        a.duration_seconds,
        a.total_questions,
        a.answered_count,
        a.correct_count,
        a.incorrect_count,
        float(a.score_percent),
        "Ha" if a.passed else ("Yo'q" if a.passed is not None else ""),
        ATTEMPT_LABELS[a.status],
    ]


def _date_format(ws: Worksheet, columns: list[int]) -> None:
    for col in columns:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for cell in row:
                cell.number_format = "DD.MM.YYYY HH:MM"


async def build_test_report_xlsx(session: AsyncSession, test_id: int) -> tuple[bytes, str]:
    test = await session.get(Test, test_id)
    if test is None:
        raise ValueError("test not found")
    summary = await participation(session, test, utcnow())
    wb = Workbook()

    ws = wb.active
    ws.title = "Umumiy"
    info = [
        ("Test", test.title),
        ("Test ID", test.id),
        ("Holat", test.status.value),
        ("Boshlanish", fmt_dt(test.starts_at)),
        ("Muddat", fmt_dt(test.deadline_at)),
        ("Savollar soni", test.question_count),
        ("O'tish foizi", test.passing_percent),
        ("Jami xodimlar", summary.total),
        ("Yakunlagan", summary.count(ParticipationStatus.COMPLETED)),
        ("Jarayonda", summary.count(ParticipationStatus.IN_PROGRESS)),
        ("Muddati o'tgan (tugatmagan)", summary.count(ParticipationStatus.EXPIRED)),
        ("Bekor qilingan", summary.count(ParticipationStatus.CANCELLED)),
        ("Qatnashmagan", summary.count(ParticipationStatus.NOT_STARTED)),
        ("Qatnashish darajasi, %", summary.participation_rate),
        ("O'rtacha ball, %", summary.average_score if summary.average_score is not None else "—"),
        ("O'tganlar", summary.passed),
        ("Hisobot yaratildi", fmt_dt(utcnow())),
    ]
    ws.append(["Ko'rsatkich", "Qiymat"])
    for row in info:
        ws.append(list(row))
    _style_sheet(ws, {1: 32, 2: 40})

    ws = wb.create_sheet("Ishtirokchilar")
    ws.append(
        [
            "Xodim",
            "Telegram ID",
            "Username",
            "Holat",
            "Foiz",
            "To'g'ri",
            "Noto'g'ri",
            "Javob berilgan",
            "Jami savollar",
            "Boshlangan",
            "Yakunlangan",
            "Davomiylik",
            "Urinishlar soni",
        ]
    )
    for r in summary.rows:
        a = r.attempt
        ws.append(
            [
                r.user.display_name,
                r.user.telegram_id,
                f"@{r.user.username}" if r.user.username else "",
                STATUS_LABELS[r.status],
                float(a.score_percent) if a and a.status != AttemptStatus.IN_PROGRESS else None,
                a.correct_count if a else None,
                a.incorrect_count if a else None,
                a.answered_count if a else None,
                a.total_questions if a else None,
                _naive_local(a.started_at) if a else None,
                _naive_local(a.completed_at) if a else None,
                fmt_duration(a.duration_seconds) if a and a.duration_seconds is not None else "",
                r.attempts_count,
            ]
        )
    _date_format(ws, [10, 11])
    _style_sheet(ws)

    ws = wb.create_sheet("Qatnashmaganlar")
    ws.append(["Xodim", "Telegram ID", "Username"])
    for r in summary.by_status(ParticipationStatus.NOT_STARTED):
        ws.append([r.user.display_name, r.user.telegram_id, f"@{r.user.username}" if r.user.username else ""])
    _style_sheet(ws)

    ws = wb.create_sheet("Barcha urinishlar")
    ws.append(ATTEMPT_HEADERS)
    rows = await session.execute(
        select(TestAttempt, User)
        .join(User, User.id == TestAttempt.user_id)
        .where(TestAttempt.test_id == test_id)
        .order_by(User.full_name, TestAttempt.attempt_no)
    )
    for attempt, user in rows:
        ws.append(_attempt_row(user, test.title, attempt))
    _date_format(ws, [6, 7])
    _style_sheet(ws)

    ws = wb.create_sheet("Xato javoblar")
    ws.append(
        [
            "Xodim",
            "Telegram ID",
            "Savol №",
            "Savol",
            "Tanlagan javob",
            "To'g'ri javob",
            "Mavzu",
            "Manba",
            "Izoh",
            "Vaqt",
        ]
    )
    rows = await session.execute(
        select(UserAnswer, TestQuestion, User)
        .join(TestQuestion, TestQuestion.id == UserAnswer.test_question_id)
        .join(User, User.id == UserAnswer.user_id)
        .where(TestQuestion.test_id == test_id, UserAnswer.is_correct.is_(False))
        .order_by(User.full_name, UserAnswer.position)
    )
    for ans, tq, user in rows:
        ws.append(
            [
                user.display_name,
                user.telegram_id,
                ans.position + 1,
                tq.question_text,
                f"{ans.selected_option}) {tq.options.get(ans.selected_option, '')}",
                f"{tq.correct_option}) {tq.options.get(tq.correct_option, '')}",
                tq.topic_name,
                tq.source_reference,
                tq.explanation,
                _naive_local(ans.answered_at),
            ]
        )
    _date_format(ws, [10])
    _style_sheet(ws, {4: 60, 5: 40, 6: 40, 8: 40, 9: 60})

    ws = wb.create_sheet("Barcha javoblar")
    ws.append(["Xodim", "Telegram ID", "Savol №", "Savol", "Tanlagan (ko'rgan harfi)", "Tanlagan variant",
               "To'g'ri (ko'rgan harfi)", "To'g'ri variant", "Natija", "Vaqt"])  # fmt: skip
    rows = await session.execute(
        select(UserAnswer, TestQuestion, User)
        .join(TestQuestion, TestQuestion.id == UserAnswer.test_question_id)
        .join(User, User.id == UserAnswer.user_id)
        .where(UserAnswer.test_id == test_id)
        .order_by(User.full_name, User.first_name, UserAnswer.position)
    )
    for ans, tq, user in rows:
        ws.append([
            user.display_name, user.telegram_id, ans.position + 1, tq.question_text, ans.selected_display,
            tq.options.get(ans.selected_option, ""), ans.correct_display or ans.correct_option,
            tq.options.get(tq.correct_option, ""), "To'g'ri" if ans.is_correct else "Xato", _naive_local(ans.answered_at),
        ])  # fmt: skip
    _date_format(ws, [10])
    _style_sheet(ws, {4: 60, 6: 40, 8: 40})

    ws = wb.create_sheet("Savollar tahlili")
    ws.append(["№", "Savol", "Mavzu", "Qiyinlik", "Javoblar", "To'g'ri", "To'g'ri %", "Manba"])
    for row in await test_question_stats(session, test_id):
        ws.append(
            [
                row.question.position,
                row.question.question_text,
                row.question.topic_name,
                row.question.difficulty.value,
                row.answers,
                row.correct,
                row.percent if row.answers else None,
                row.question.source_reference,
            ]
        )
    _style_sheet(ws, {2: 70, 8: 40})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"test_{test.id}_hisobot_{to_local(utcnow()):%Y%m%d_%H%M}.xlsx"


async def build_summary_report_xlsx(session: AsyncSession, since: datetime | None = None) -> tuple[bytes, str]:
    """All attempts (optionally since a date) + per-employee summary sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    ws.append(ATTEMPT_HEADERS)
    stmt = (
        select(TestAttempt, User, Test.title)
        .join(User, User.id == TestAttempt.user_id)
        .join(Test, Test.id == TestAttempt.test_id)
        .order_by(TestAttempt.started_at.desc())
    )
    if since is not None:
        stmt = stmt.where(TestAttempt.started_at >= since)
    per_user: dict[int, dict] = {}
    for attempt, user, title in await session.execute(stmt):
        ws.append(_attempt_row(user, title, attempt))
        if user.role != UserRole.EMPLOYEE or attempt.status == AttemptStatus.CANCELLED:
            continue
        agg = per_user.setdefault(
            user.id,
            {
                "user": user,
                "attempts": 0,
                "completed": 0,
                "expired": 0,
                "questions": 0,
                "correct": 0,
                "incorrect": 0,
                "seconds": 0,
                "scores": [],
            },
        )
        if attempt.status == AttemptStatus.IN_PROGRESS:
            continue
        agg["attempts"] += 1
        agg["completed"] += attempt.status == AttemptStatus.COMPLETED
        agg["expired"] += attempt.status == AttemptStatus.EXPIRED
        agg["questions"] += attempt.total_questions
        agg["correct"] += attempt.correct_count
        agg["incorrect"] += attempt.incorrect_count
        agg["seconds"] += attempt.duration_seconds or 0
        agg["scores"].append(float(attempt.score_percent))
    _date_format(ws, [6, 7])
    _style_sheet(ws)

    ws = wb.create_sheet("Xodimlar bo'yicha")
    ws.append(
        [
            "Xodim",
            "Telegram ID",
            "Username",
            "Urinishlar",
            "Yakunlangan",
            "Muddati o'tgan",
            "Jami savollar",
            "To'g'ri",
            "Noto'g'ri",
            "Umumiy aniqlik %",
            "O'rtacha ball %",
            "Eng yaxshi %",
            "Eng past %",
            "Jami vaqt",
        ]
    )
    for agg in sorted(per_user.values(), key=lambda a: -(a["correct"] / a["questions"]) if a["questions"] else 0):
        scores = agg["scores"]
        user = agg["user"]
        ws.append(
            [
                user.display_name,
                user.telegram_id,
                f"@{user.username}" if user.username else "",
                agg["attempts"],
                agg["completed"],
                agg["expired"],
                agg["questions"],
                agg["correct"],
                agg["incorrect"],
                round(agg["correct"] / agg["questions"] * 100, 1) if agg["questions"] else None,
                round(sum(scores) / len(scores), 1) if scores else None,
                max(scores) if scores else None,
                min(scores) if scores else None,
                fmt_duration(agg["seconds"]),
            ]
        )
    _style_sheet(ws)

    ws = wb.create_sheet("Testlar")
    ws.append(["ID", "Test", "Holat", "Boshlanish", "Muddat", "Savollar", "O'tish foizi"])
    tests = await session.execute(select(Test).where(Test.status != TestStatus.DRAFT).order_by(Test.starts_at.desc()))
    for test in tests.scalars():
        ws.append(
            [
                test.id,
                test.title,
                test.status.value,
                _naive_local(test.starts_at),
                _naive_local(test.deadline_at),
                test.question_count,
                test.passing_percent,
            ]
        )
    _date_format(ws, [4, 5])
    _style_sheet(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"umumiy_hisobot_{to_local(utcnow()):%Y%m%d_%H%M}.xlsx"


async def build_test_report_csv(session: AsyncSession, test_id: int) -> tuple[bytes, str]:
    test = await session.get(Test, test_id)
    if test is None:
        raise ValueError("test not found")
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(ATTEMPT_HEADERS)
    rows = await session.execute(
        select(TestAttempt, User)
        .join(User, User.id == TestAttempt.user_id)
        .where(TestAttempt.test_id == test_id)
        .order_by(User.full_name, TestAttempt.attempt_no)
    )
    for attempt, user in rows:
        row = _attempt_row(user, test.title, attempt)
        row[5] = fmt_dt(attempt.started_at)
        row[6] = fmt_dt(attempt.completed_at) if attempt.completed_at else ""
        writer.writerow(row)
    # UTF-8 BOM so Excel opens Cyrillic/Uzbek text correctly.
    return ("﻿" + buf.getvalue()).encode("utf-8"), f"test_{test.id}_natijalar.csv"
